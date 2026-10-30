# カード資料(2026-10-30)を作る: CardCatalogExport の TSV + 翻訳の表 + roles.py → xlsx / md
#   python catalog.py <TSVのフォルダ> <出力フォルダ>
# (2026-10-30 属性20枚の再設計の式に合わせた。文章の所見 A(...) は再設計前の資料のもの)
# 数値の計算は ElementStats / CardRules.Scale / GachaStage / GameManager.DrawFromGachaPool と同じ式(式を変えた時はここも合わせる)
import csv, json, math, os, re, sys
from collections import Counter, OrderedDict
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from roles import R

SRC, OUT = sys.argv[1], sys.argv[2]
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
LOC = os.path.join(ROOT, "Unity", "2ndAction", "Assets", "Resources", "Localization")
os.makedirs(OUT, exist_ok=True)

def tsv(name):
    with open(os.path.join(SRC, name), encoding="utf-8") as f:
        return list(csv.DictReader(f, delimiter="\t", quoting=csv.QUOTE_NONE))

cards = tsv("cards.tsv")
combos = tsv("combos.tsv")
settings = {r["key"]: r["value"] for r in tsv("settings.tsv")}
en = json.load(open(os.path.join(LOC, "en.json"), encoding="utf-8")); en = dict(zip(en["k"], en["v"]))
ja = json.load(open(os.path.join(LOC, "ja.json"), encoding="utf-8")); ja = dict(zip(ja["k"], ja["v"]))

CAT_JA = OrderedDict([("Attack", "ATTACK"), ("Defense", "DEFENSE"), ("Movement", "MOVEMENT"), ("Growth", "GROWTH"), ("Heal", "HEAL"), ("Special", "SPECIAL"), ("Risk", "RISK")])
EL_JA = OrderedDict([("Fire", "炎"), ("Ice", "氷"), ("Thunder", "雷"), ("Wind", "風"), ("Blood", "血(その他)"), ("None", "なし")])
W = [float(x) for x in settings["gachaRarityWeights"].split(",")]
TH = [float(x) for x in settings["gachaThresholds"].split(",")]
assert len(cards) == 100, len(cards)
byid = {c["id"]: c for c in cards}
for c in cards:
    assert c["id"] in R, c["id"]

# ---------------------------------------------------------------- 抽選
def gacha_probs(dist, stage):
    pool = [c for c in cards if dist >= float(c["unlockDistance"]) and stage >= int(c["gachaStage"])]
    buckets = Counter(int(c["rarity"]) for c in pool)
    tw = sum(W[r - 1] for r in buckets)
    p = {}
    for c in pool:
        r = int(c["rarity"])
        p[c["id"]] = W[r - 1] / tw / buckets[r]
    return p, pool, buckets, tw
STAGES = [(1, 0.0, "段階1 OLD (0m〜)"), (1, 500.0, "段階1 (BEST 500m〜)"), (2, 5000.0, "段階2 REPAIRED (5,000m〜)"), (3, 20000.0, "段階3 MECHANICAL (20,000m〜)"), (4, 50000.0, "段階4 MAGICAL (50,000m〜)"), (5, 100000.0, "段階5 DEATH-TOUCHED (100,000m〜)")]
stage_probs = [(lab,) + gacha_probs(d, s) for s, d, lab in STAGES]
p_full = stage_probs[-1][1]

def unlock_text(c):
    d = float(c["unlockDistance"]); s = int(c["gachaStage"])
    if d <= 0 and s <= 1: return "最初から(ガチャ段階1)"
    need = max(d, TH[s - 1])
    return f"BEST {need:,.0f}m 以上(ガチャ段階{s})"

# ---------------------------------------------------------------- 属性の計算(ElementStats と同じ)
def scale(kind, v, lv):
    if lv <= 0: return 0.0
    if kind == "Once": return v
    if kind == "Every3": return v * math.ceil(lv / 3)
    if kind == "Every2": return v * math.ceil(lv / 2)
    if kind == "Phoenix": return v * (1 + lv // 2)
    if kind == "After3": return v * ((lv - 1) // 3)
    return v * lv
def effects(c):
    out = []
    for part in c["effects"].split(" ; "):
        if not part: continue
        t, v, k = part.split(":")
        out.append((t, float(v), k))
    return out
def total(levels):
    s = Counter()
    for cid, lv in levels.items():
        for t, v, k in effects(byid[cid]): s[t] += scale(k, v, lv)
    return s

def el_stats(el, levels):
    g = total(levels)
    clamp = lambda x, a, b: max(a, min(b, x))
    if el == "Fire":
        on = g["BurnChance"] > 0 and g["BurnPower"] > 0
        dps = g["BurnPower"] * (1 + max(0, g["BurnPowerPct"])); dur = 3 + max(0, g["BurnDuration"])
        s = f"炎上 確率{min(1,g['BurnChance'])*100:.0f}% / 毎秒{dps:.1f} × {dur:.2g}秒(1回の炎上で約{dps*dur:.0f}、重ねると最大+50%・ボスは+15%)" if on else "炎上しない"
        if g["InfernoLevel"] > 0:
            lv = int(g["InfernoLevel"]); s += f" / 延焼(炎上中の敵を撃破) 威力{8+2.5*lv:g}・半径{2.2+0.12*lv:.2f}m・{2+(lv-1)//3}段"
        if g["FlameCounterLevel"] > 0:
            lv = int(g["FlameCounterLevel"]); s += f" / Shieldで防ぐと炎の反撃 威力{round(20*(1.0+0.2*lv))}"
        return on, s
    if el == "Ice":
        slow = clamp(g["ChillSlow"], 0, 0.7)
        on = g["ChillChance"] > 0 and slow > 0
        fz = max(2, 5 - round(g["FreezeThresholdReduce"])) if slow > 0 else 0
        s = f"冷え 確率{min(1,g['ChillChance'])*100:.0f}% / 行動-{slow*100:.0f}% × {2+max(0,g['ChillDuration']):.2g}秒 / 凍結={fz}回冷やす → {1+max(0,g['FreezeDuration']):.2g}秒停止(ボスは強い減速)"
        if g["AbsoluteZeroLevel"] > 0:
            lv = int(g["AbsoluteZeroLevel"]); s += f" / 冷えた敵へ攻撃+{3*lv}%・凍結で砕け 威力{12+4*lv}・ボスのBREAK蓄積+{6*lv}%"
        return on, s
    if el == "Thunder":
        on = g["LightningChance"] > 0 and g["LightningPower"] > 0
        dmg = g["LightningPower"] * (1 + max(0, g["LightningDamagePct"])) * 0.5 # ElementStats.LightningScale(2026-10-30)
        s = f"落雷 確率{min(1,g['LightningChance'])*100:.1f}% / 威力{dmg:.1f} / 連鎖{int(clamp(round(g['LightningChains']),0,8))}体(距離{3.5+g['LightningRange']:.2g}m) / 範囲{clamp(g['LightningSplash'],0,3):.2g}m / 雑魚停止{clamp(g['LightningShock'],0,0.6):.2g}秒"
        return on, s
    if el == "Wind":
        on = (g["WindBladeChance"] > 0 and g["WindBladePower"] > 0)
        s = f"風刃 確率{min(1,g['WindBladeChance'])*100:.0f}% / 威力{g['WindBladePower']:.1f} / 貫通{2+int(round(g['WindPierce']))}体 / 射程{6+g['WindRange']:.2g}m / 速さ×{1+g['WindSpeedPct']:.2f}"
        if g["TornadoLevel"] > 0:
            lv = int(g["TornadoLevel"]); s += f" / 風刃の命中で竜巻 威力{5+1.25*lv:g}×{3+(lv-1)//3}・半径{1.4+0.1*lv:.2f}m・引き寄せ{0.25+0.04*lv:.2f}m(ボスへは半分)"
        return on, s
    if el == "Blood":
        on = g["BleedChance"] > 0 and g["BleedPower"] > 0
        parts = []
        if g["LifestealChance"] > 0: parts.append(f"吸収 確率{g['LifestealChance']*100:.0f}%")
        if g["LowHp25LifestealChance"] > 0: parts.append(f"HP25%以下で吸収+{g['LowHp25LifestealChance']*100:.0f}%")
        if g["PredatorLevel"] > 0: parts.append("ボスにも吸収(1/10)")
        if g["OverhealLevel"] > 0: parts.append("溢れた回復→Shield")
        if g["HealBonusHearts"] > 0: parts.append(f"回復量+ハート{g['HealBonusHearts']:.0f}")
        if on: parts.append(f"出血 確率{min(1,g['BleedChance'])*100:.0f}%(満HPで×2)/ 毎秒{g['BleedPower']:.1f} × {4+g['BleedDuration']:.2g}秒")
        if g["PhoenixLevel"] > 0: parts.append(f"復活 ハート{1+int(levels.get('phoenix',0))//2}")
        if g["LowHp50Pct"] > 0: parts.append(f"HP50%以下 攻撃+{g['LowHp50Pct']*100:.0f}%")
        return on or bool(parts), " / ".join(parts)
    return False, ""

# 属性ごとのカード(並びは系統の流れ)
FAMILY = OrderedDict([
    ("Fire", ["flame_blade", "burning_soul", "inferno", "flame_counter"]),
    ("Ice", ["frost_edge", "ice_prison", "absolute_zero"]),
    ("Thunder", ["thunder_strike", "chain_lightning", "high_voltage", "thunder_lord"]),
    ("Wind", ["wind_cutter", "gale", "tornado"]),
    ("Blood", ["vampire", "predator", "overheal", "blood_rush", "blood_blade", "phoenix"]),
])
for el, ids in FAMILY.items():
    for i in ids: assert byid[i]["element"] == el, (i, byid[i]["element"])
assert sum(len(v) for v in FAMILY.values()) == sum(1 for c in cards if c["element"] != "None")
COMBO_OF = {}
for k in combos:
    for a in k["abilities"].split("+"): COMBO_OF.setdefault(a, []).append(k["name"])

# ---------------------------------------------------------------- 1行ずつ
def row(c):
    cid = c["id"]; use, tags, pre = R[cid]
    n = c["name"]
    return OrderedDict([
        ("No", int(c["no"])), ("ID", cid), ("日本語名", ja.get(n, n) + ("" if ja.get(n, n) != n else "(英語名のまま)")), ("英語名", en.get(n, n)),
        ("レア度", "★" * int(c["rarity"])), ("★", int(c["rarity"])), ("カテゴリ", CAT_JA[c["category"]]), ("最大Lv", int(c["maxLv"])),
        ("Lv1の効果", c["lv1"]), ("Lv9の効果", c["lv9"]), ("説明文(日本語)", c["description"]), ("説明文(英語)", en.get(c["description"], "")),
        ("主な用途", use), ("役割タグ", tags), ("属性", "あり" if c["element"] != "None" else "なし"), ("属性の種類", EL_JA[c["element"]]),
        ("属性1枚で発動(Lv1)", c["lv1Element"] or ("—" if c["element"] == "None" else "単独では属性の効果なし")),
        ("COMBO対象", "、".join(COMBO_OF.get(cid, [])) or "—"),
        ("FINAL EVOLUTION", "対象外" if not c["feTitle"] else f"{c['feTitle']}: {c['feDesc']}"),
        ("出現条件", "所持してデッキに入れると、ラン中の LEVEL UP / ボス報酬の3択に出る" + ("(マルチでは出ない)" if cid == "character_ultimate" else "") + ("(HP封印を払えない時は次のLvが出ない)" if int(c["sacrificeLv9"]) > 0 else "")),
        ("解放条件(ガチャ)", unlock_text(c)), ("ガチャ段階", int(c["gachaStage"])), ("解放距離(m)", float(c["unlockDistance"])),
        ("ガチャ1回で出る確率(全解放時)", p_full.get(cid, 0.0)), ("おすすめ編成の優先度", int(c["recommendPriority"])),
        ("HP封印 Lv1/Lv9", f"{c['sacrificeLv1']}/{c['sacrificeLv9']}" if int(c["sacrificeLv9"]) > 0 else "—"),
        ("特殊な前提", pre or "—"), ("内部の効果(種類:値:伸び方)", c["effects"]),
    ])
rows = [row(c) for c in cards]

# ---------------------------------------------------------------- 集計
rar = Counter(int(c["rarity"]) for c in cards)
cat_r = {k: Counter() for k in CAT_JA}
for c in cards: cat_r[c["category"]][int(c["rarity"])] += 1
el_r = {k: Counter() for k in EL_JA}
for c in cards: el_r[c["element"]][int(c["rarity"])] += 1
el_cat = {k: Counter() for k in EL_JA}
for c in cards: el_cat[c["element"]][c["category"]] += 1
tag_n = Counter(t for c in cards for t in R[c["id"]][1].split(","))

# ---------------------------------------------------------------- xlsx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
wb = Workbook()
HF = Font(bold=True, color="FFFFFF"); HFill = PatternFill("solid", fgColor="1F3A5F")
thin = Side(style="thin", color="C8CDD5")
def sheet(title, header, data, widths=None, wrap=True, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for r in data: ws.append(list(r))
    for cell in ws[1]: cell.font = HF; cell.fill = HFill; cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row_ in ws.iter_rows(min_row=2):
        for cell in row_:
            cell.alignment = Alignment(vertical="top", wrap_text=wrap)
            cell.border = Border(bottom=thin)
    if widths:
        for i, w in enumerate(widths): ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.freeze_panes = "C2" if len(header) > 6 else "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws

hdr = list(rows[0].keys())
ws = sheet("全100枚", hdr, [list(r.values()) for r in rows],
           [5, 18, 20, 18, 9, 4, 11, 6, 34, 34, 44, 44, 26, 18, 6, 9, 16, 22, 40, 34, 26, 7, 9, 10, 8, 9, 26, 40], first=True)
for cell in ws["X"][1:]: cell.number_format = "0.00%"
ELFILL = {"炎": "FDE2D2", "氷": "DCEFFC", "雷": "FFF6CC", "風": "DDF6E8", "血(その他)": "F6D8DD"}
for r in ws.iter_rows(min_row=2):
    f = ELFILL.get(r[15].value)
    if f: r[15].fill = PatternFill("solid", fgColor=f)

# 属性
el_rows = []
for el, ids in FAMILY.items():
    for i in ids:
        c = byid[i]
        el_rows.append([EL_JA[el], "★" * int(c["rarity"]), i, c["name"], CAT_JA[c["category"]], c["lv1"], c["lv9"], c["lv1Element"] or "単独では属性の効果なし", "、".join(COMBO_OF.get(i, [])) or "—", c["feTitle"] + ": " + c["feDesc"] if c["feTitle"] else "—", unlock_text(c)])
sheet("属性カード", ["属性", "レア度", "ID", "名前", "カテゴリ", "Lv1の効果", "Lv9の効果", "1枚で起きること(Lv1)", "COMBO", "FINAL EVOLUTION", "解放条件"], el_rows, [9, 8, 16, 16, 10, 40, 40, 22, 24, 40, 24])

def scenarios(el, ids):
    base = ids[0]
    sc = [(f"1枚: {byid[base]['name']} Lv1", {base: 1}), (f"1枚: {byid[base]['name']} Lv9", {base: 9})]
    for i in ids[1:]:
        sc.append((f"1枚: {byid[i]['name']} Lv1(入口のカード無し)", {i: 1}))
        sc.append((f"1枚: {byid[i]['name']} Lv9(入口のカード無し)", {i: 9}))
    sc.append(("系統を全部 Lv1", {i: 1 for i in ids}))
    sc.append(("系統を全部 Lv9(上位カードまで揃えた)", {i: 9 for i in ids}))
    return [(lab, el_stats(el, lv)) for lab, lv in sc]
sc_rows = []
for el, ids in FAMILY.items():
    for lab, (on, s) in scenarios(el, ids): sc_rows.append([EL_JA[el], lab, "発動" if on else "発動しない", s])
sheet("属性の段階", ["属性", "場合", "属性の効果", "数値(攻撃の強化なし・1体あたり。通常の1発 = 20)"], sc_rows, [9, 40, 10, 110])

dist = [["合計"] + [rar[i] for i in range(1, 6)] + [sum(rar.values())]]
for k, v in CAT_JA.items(): dist.append([v] + [cat_r[k][i] for i in range(1, 6)] + [sum(cat_r[k].values())])
dist.append([]); dist.append(["属性別"] + ["★1", "★2", "★3", "★4", "★5", "計"])
for k, v in EL_JA.items(): dist.append([v] + [el_r[k][i] for i in range(1, 6)] + [sum(el_r[k].values())])
dist.append([]); dist.append(["属性×カテゴリ"] + list(CAT_JA.values()))
for k, v in EL_JA.items(): dist.append([v] + [el_cat[k][c] for c in CAT_JA])
dist.append([]); dist.append(["役割タグ(アイリーンの判断)", "枚数"])
for t, n in tag_n.most_common(): dist.append([t, n])
sheet("レア度分布", ["", "★1", "★2", "★3", "★4", "★5", "計"], dist, [26, 8, 8, 8, 8, 8, 8], wrap=False)

sheet("COMBO", ["ID", "名前", "必要な2枚", "効果", "最終進化中(ENHANCED)", "リスク"],
      [[k["id"], k["name"], " + ".join(byid[a]["name"] for a in k["abilities"].split("+")), k["description"], k["enhanced"], "あり" if k["risk"] == "1" else ""] for k in combos], [16, 16, 30, 60, 40, 6])
sheet("FINAL EVOLUTION", ["No", "ID", "名前", "★", "題", "効果", "種類", "長さ(現行は全カード60秒)", "自分の効果の倍率", "追加の効果"],
      [[int(c["no"]), c["id"], c["name"], int(c["rarity"]), c["feTitle"], c["feDesc"], c["feModule"], c["feDuration"], c["feAmplify"], c["feBonuses"]] for c in cards if c["feTitle"]],
      [5, 18, 18, 4, 14, 60, 12, 10, 8, 40])

gs = [["ガチャのレア度の重み", "★1 50 / ★2 30 / ★3 15 / ★4 4 / ★5 1(引ける★の中だけで割る)"], ["", "まず★を重みで選び、その★の中で等確率に1枚"], []]
gs.append(["段階", "引ける枚数", "★1", "★2", "★3", "★4", "★5", "1枚あたり ★1", "★2", "★3", "★4", "★5"])
for lab, p, pool, buckets, tw in stage_probs:
    per = [(W[r - 1] / tw / buckets[r]) if buckets[r] else 0 for r in range(1, 6)]
    gs.append([lab, len(pool)] + [buckets[r] for r in range(1, 6)] + [f"{x*100:.2f}%" if x else "—" for x in per])
gs += [[], ["ラン中の LEVEL UP / ボス報酬", "デッキ(最大12枚)の中から、Lv9に届いていないカードを等確率で混ぜて3枚(同じカードは1回)。レア度の重みは無い"],
       ["", "FINAL EVOLUTION の候補(READY)があれば3枠のうち最大1枠(候補が足りない時はFEで埋める)"],
       ["", "デッキが全部Lv9なら3択は出ない。デッキが空なら出ない(デッキが壊れた時だけ解放済みの全カードから)"],
       ["", "HP封印のカードは、次のLvの封印を払えない時は候補に出ない。#100 ULTIMATE はマルチでは出ない"],
       ["", "デッキのLv2以上のカード(合成した物)を選ぶと、その強化量ぶん一度に上がる(能力ごとに Lv9 まで)"],
       ["キャラカード", "キャラごとに3枠。ランの開始時にそのLvで付く(デッキとは別)"],
       ["COMBO", "必要な2枚をそのランで両方持つと成立(Lv1+Lv1でも)。マルチでは効果なし"],
       ["FINAL EVOLUTION", f"そのランで能力Lv9 → {float(settings['feReadyMeters']):,.0f}m 走ると READY。選ぶと60秒(実際に進んでいる時間)だけ限界突破。#100 ULTIMATE は対象外。マルチでは出ない"]]
sheet("抽選と出現", ["項目", "内容"], gs, [30, 14, 6, 6, 6, 6, 6, 12, 8, 8, 8, 8], wrap=False)

# 段階ごとの1枚あたりの確率(カードごと)
st_rows = []
for c in cards:
    st_rows.append([int(c["no"]), c["id"], c["name"], int(c["rarity"])] + [sp[1].get(c["id"], 0.0) for sp in stage_probs])
ws2 = sheet("ガチャ確率(カード別)", ["No", "ID", "名前", "★"] + [sp[0] for sp in stage_probs], st_rows, [5, 18, 18, 4, 14, 14, 14, 14, 14, 14])
for col in "EFGHIJ":
    for cell in ws2[col][1:]: cell.number_format = "0.00%"

guide = [["この資料", "2026-10-30 時点の実装(コード/アセット)から自動で書き出した値。性能・レア度・出現率は変えていない"],
         ["Lv1/Lv9の効果", "ゲーム内の詳細表示と同じ計算(CardVariant.AbilityEffectTextAt)。ATTACK UP は高いLvほど伸びが大きい(BalanceTuning)"],
         ["最大Lv", "全カード Lv9(所持の合成Lv・ラン中の能力Lv とも)。Lv9 の後は Mastery ★1〜5 → AWAKENED(性能は上がらない)"],
         ["日本語名", "カード名は日本語の画面でも英語名のまま(日本語名は未設定)。説明文は30言語"],
         ["カテゴリ", "CardCategory の7種: ATTACK / DEFENSE / MOVEMENT / GROWTH / HEAL / SPECIAL / RISK"],
         ["属性", "CardDefinition.element: 炎/氷/雷/風/血(その他)。属性の効果は ElementSystem(属性同士の相性・相互作用は無い)"],
         ["役割タグ・主な用途・特殊な前提", "アイリーンの判断(相談の叩き台)。Tools/cards/roles.py"],
         ["作り直し", "Unity: -executeMethod CardCatalogExport.Run -ccOut <dir> → python Tools/cards/catalog.py <dir> <出力>"]]
sheet("読み方", ["項目", "内容"], guide, [26, 110])
wb.move_sheet("読み方", offset=-(len(wb.sheetnames) - 1))
xlsx = os.path.join(OUT, "card_catalog.xlsx"); wb.save(xlsx)

# ---------------------------------------------------------------- md
L = []
A = L.append
A("# One More Mile 現行カード 100枚 資料(2026-10-30)\n")
A("実装(コード/アセット)から自動で書き出した値です。**レア度・性能・属性・出現率は変更していません。** 全項目の一覧は `card_catalog.xlsx`(シート「全100枚」)。\n")
A("## 1. レア度の分布\n")
A("| | ★1 | ★2 | ★3 | ★4 | ★5 | 計 |\n|---|---|---|---|---|---|---|")
for r in dist[:1 + len(CAT_JA)]: A("| " + " | ".join(str(x) for x in r) + " |")
A("\n**属性別**\n\n| 属性 | ★1 | ★2 | ★3 | ★4 | ★5 | 計 |\n|---|---|---|---|---|---|---|")
for k, v in EL_JA.items(): A(f"| {v} | " + " | ".join(str(el_r[k][i]) for i in range(1, 6)) + f" | {sum(el_r[k].values())} |")
A("\n**属性 × カテゴリ**\n\n| 属性 | " + " | ".join(CAT_JA.values()) + " |\n" + "|---" * (len(CAT_JA) + 1) + "|")
for k, v in EL_JA.items(): A(f"| {v} | " + " | ".join(str(el_cat[k][c]) for c in CAT_JA) + " |")
A("\n## 2. 属性カード\n")
A("属性は「命中した時に確率で起きる追加の現象」です。どのキャラでも同じカードなら同じ効果になります。属性同士の相性や相互作用はありません(炎で凍結が溶ける、なども無し)。")
A("属性のダメージは、属性ごとの基礎の値に **(1 + 無条件の攻撃の強化の3/4)**(CardRules.ElementAttackShare 0.75。2026-10-30 の再設計で 0.5 → 0.75) を掛けたものです。条件付きの攻撃カード(空中・締め・ボス特効など)は乗りません。下の数値は攻撃の強化が無い時の値で、通常の1発は 20(黒剣士の基礎)です。\n")
TREE = {
 "Fire": "FLAME BLADE ★4(入口: 命中で炎上)\n ├─ BURNING SOUL ★3(炎上を強く・長く) ── COMBO BLAZING EDGE(FLAME BLADE+BURNING SOUL)\n │    └─ INFERNO ★4(炎上中の敵を倒すと延焼) ── COMBO WILDFIRE(BURNING SOUL+INFERNO)\n └─ FLAME COUNTER ★3(DEFENSE: Shieldで防ぐと炎の反撃。系統のCOMBOは無し)",
 "Ice": "FROST EDGE ★4(入口: 命中で冷やす)\n └─ ICE PRISON ★3(凍結しやすく) ── COMBO FROZEN PRISON(FROST EDGE+ICE PRISON)\n      └─ ABSOLUTE ZERO ★4(冷えた敵へ火力+凍結で砕け) ── COMBO SHATTER(ICE PRISON+ABSOLUTE ZERO)",
 "Thunder": "THUNDER STRIKE ★4(入口: 命中で落雷)\n ├─ CHAIN LIGHTNING ★3(連鎖) ── COMBO THUNDER CHAIN(THUNDER STRIKE+CHAIN LIGHTNING)\n │    └─ THUNDER LORD ★4(範囲・連鎖・威力) ── COMBO STORM LORD(CHAIN LIGHTNING+THUNDER LORD)\n └─ HIGH VOLTAGE ★4(威力・停止。COMBOは無し)",
 "Wind": "WIND CUTTER ★4(入口: 命中で風刃)\n └─ GALE ★3(風刃の速さ・射程・貫通) ── COMBO GALE EDGE(WIND CUTTER+GALE)\n      └─ TORNADO ★4(風刃の命中で竜巻) ── COMBO CYCLONE(GALE+TORNADO)",
 "Blood": "VAMPIRE ★3(吸収) ── COMBO BLOOD AEGIS(VAMPIRE+OVERHEAL)\n ├─ PREDATOR ★3(ボスにも吸収)\n ├─ OVERHEAL ★3(溢れた回復→Shield)\n ├─ BLOOD RUSH ★3(低HPで火力と吸収)\nBLOOD BLADE ★4(出血。血の系統で唯一の命中時の継続ダメージ)\nPHOENIX ★5(復活)\n※ 血は「吸収/回復/低HP」の寄せ集めで、炎〜風のような入口→上位の1本の流れにはなっていない",
}
for el, ids in FAMILY.items():
    n = len(ids); rr = Counter(int(byid[i]["rarity"]) for i in ids)
    A(f"### {EL_JA[el]}({n}枚: " + " / ".join(f"★{r} {rr[r]}" for r in sorted(rr)) + ")\n")
    RAR_BY_NAME = {c["name"]: c["rarity"] for c in cards}
    A("```\n" + re.sub(r"([A-Z][A-Z ]*[A-Z]) ★\d", lambda m: m.group(1) + " ★" + RAR_BY_NAME.get(m.group(1), "?"), TREE[el]) + "\n```\n") # ★は今の値で
    A("| 場合 | 属性の効果 | 数値 |\n|---|---|---|")
    for lab, (on, s) in scenarios(el, ids): A(f"| {lab} | {'発動' if on else '発動しない'} | {s} |")
    A("")
A("### 属性の気づき\n")
A("- 2026-10-30 の属性20枚の再設計(役割・レア度・解放・Lv1/Lv9・攻撃の強化の乗り方)の所見と試験の数値は Docs/ElementBalance_2026-10-30/report.md。再設計前の所見は Docs/CardCatalog_2026-10-30/card_catalog.md。\n")
A("## 3. COMBO 対象(15組・30枚が関係)\n")
A("| COMBO | 必要な2枚 | 効果 |\n|---|---|---|")
for k in combos: A(f"| {k['name']}{'(リスク)' if k['risk']=='1' else ''} | " + " + ".join(byid[a]["name"] for a in k["abilities"].split("+")) + f" | {k['description']} |")
in_combo = sorted({a for k in combos for a in k["abilities"].split("+")}, key=lambda i: int(byid[i]["no"]))
A(f"\nCOMBO に関係するカード: {len(in_combo)}枚 — " + "、".join(byid[i]["name"] for i in in_combo) + "\n")
A("## 4. FINAL EVOLUTION 対象\n")
fe_n = sum(1 for c in cards if c["feTitle"])
A(f"**{fe_n}枚**(100枚のうち #100 ULTIMATE だけ対象外)。そのランで能力 Lv9 → {float(settings['feReadyMeters']):,.0f}m 走ると READY → LEVEL UP / ボス報酬の3択に最大1枠 → 選ぶと60秒(ランが実際に進んでいる時間)だけ限界突破。終わるとすぐ READY に戻ります。専用の処理がある代表カードは ATTACK UP / SPEED UP / ATTACK RANGE UP / VAMPIRE / PHOENIX / FLAME BLADE / THUNDER STRIKE / EXP UP / GREED。ほかは「自分の効果×倍率 + 追加の効果」です。各カードの内容は xlsx の「FINAL EVOLUTION」シート。\n")
A("## 5. 出現条件と抽選の仕様(現行)\n")
A("1. **解放(ガチャに出るようになる)**: BEST 距離とガチャ段階の両方を満たしたカード。段階は BEST 0 / 5,000 / 20,000 / 50,000 / 100,000m。PATHFINDER だけ BEST 500m。")
A("2. **ガチャ**: まず★を重み(★1 50 / ★2 30 / ★3 15 / ★4 4 / ★5 1、引ける★の中だけで割る)で選び、その★の中で等確率に1枚。")
A("3. **ラン中**: デッキ(最大12枚)から、Lv9 未満のカードを等確率で混ぜて3枚(同じカードは1回)。**レア度による重みはありません。** FINAL EVOLUTION の候補があれば最大1枠。")
A("4. **キャラカード**: キャラごとに3枠。ランの開始時にそのLvで付きます(デッキとは別)。\n")
A("| ガチャの段階 | 引ける枚数 | ★1 | ★2 | ★3 | ★4 | ★5 | 1枚あたりの確率 ★1 / ★2 / ★3 / ★4 / ★5 |\n|---|---|---|---|---|---|---|---|")
for lab, p, pool, buckets, tw in stage_probs:
    per = [(W[r - 1] / tw / buckets[r]) if buckets[r] else 0 for r in range(1, 6)]
    A(f"| {lab} | {len(pool)} | " + " | ".join(str(buckets[r]) for r in range(1, 6)) + " | " + " / ".join(f"{x*100:.2f}%" if x else "—" for x in per) + " |")
A("\n**抽選の気づき(相談用)**\n")
_s3 = next(x for x in stage_probs if x[0].startswith("段階3")); _s5 = stage_probs[-1]
_pc = lambda sp, r: (W[r - 1] / sp[4] / sp[3][r]) if sp[3][r] else 0
A(f"- 段階3(20,000m)の★5は {_s3[3][5]}枚で、1枚あたり **{_pc(_s3, 5)*100:.2f}%**(★4 の1枚 {_pc(_s3, 4)*100:.2f}% の約{_pc(_s3, 5)/_pc(_s3, 4):.0f}倍)。")
A(f"- 全解放後の1枚あたりは ★4 = {_pc(_s5, 4)*100:.2f}%、★5 = {_pc(_s5, 5)*100:.2f}%、★3 = {_pc(_s5, 3)*100:.2f}%、★1 = {_pc(_s5, 1)*100:.2f}%。")
A("- ラン中の3択はレア度の重みが無く、デッキに入れた12枚が等確率です。レア度は「ガチャで出にくい」だけで、ラン中の出やすさには関係しません。")
A("\n## 6. 全100枚(要約。全項目は xlsx)\n")
A("| No | 名前 | ★ | カテゴリ | 属性 | Lv1 | Lv9 | 主な用途 | 役割タグ | COMBO | 解放 |\n|---|---|---|---|---|---|---|---|---|---|---|")
for r in rows:
    A(f"| {r['No']} | {r['英語名']} | {r['★']} | {r['カテゴリ']} | {r['属性の種類'] if r['属性']=='あり' else ''} | {r['Lv1の効果']} | {r['Lv9の効果']} | {r['主な用途']} | {r['役割タグ']} | {r['COMBO対象'] if r['COMBO対象']!='—' else ''} | {r['解放条件(ガチャ)']} |")
A("\n## 7. 役割タグの数(アイリーンの判断)\n")
A(" / ".join(f"{t} {n}" for t, n in tag_n.most_common()))
md = os.path.join(OUT, "card_catalog.md")
open(md, "w", encoding="utf-8").write("\n".join(L) + "\n")
print("ok", xlsx, md, dict(rar), {k: dict(v) for k, v in el_r.items()})
